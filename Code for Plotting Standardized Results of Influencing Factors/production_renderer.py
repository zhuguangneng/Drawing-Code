#!/usr/bin/env python
"""Create the bilingual two-level factor trees and auditable frequency data.

The formal source workbook is always read-only.  Workbook authoring is delegated
to the companion artifact-tool builder; this script prepares the statistics and
static vector/raster figures, then performs final cross-format validation.
"""

from __future__ import annotations

import argparse
import hashlib
import html
import json
import math
import re
import subprocess
import sys
import unicodedata
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable
import xml.etree.ElementTree as ET

import pandas as pd
from PIL import Image, ImageDraw, ImageFont, ImageStat
from pypdf import PdfReader
from reportlab.lib.colors import HexColor
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas


ROOT = Path(r"C:\Users\LENOVO\Desktop\因素归一标准化")
INPUT = ROOT / "output" / "因素归类标准化_结果.xlsx"
CODEBOOK = ROOT / "output" / "标准因素与一级分类代码本.xlsx"
OUT = ROOT / "output" / "visualization"
WORK = ROOT / "scripts" / "visualization_work"
DATA_JSON = WORK / "visualization_data.json"
ARTIFACT_VISUAL_JSON = WORK / "artifact_visual_validation.json"
VALIDATION_JSON = OUT / "可视化验证报告.json"
LOG_PATH = OUT / "可视化运行日志.md"

CN_SVG = OUT / "01_标准因素两级树状图_中文.svg"
CN_PDF = OUT / "01_标准因素两级树状图_中文.pdf"
CN_PNG = OUT / "01_标准因素两级树状图_中文.png"
EN_SVG = OUT / "02_Standard_Factor_Two_Level_Tree_EN.svg"
EN_PDF = OUT / "02_Standard_Factor_Two_Level_Tree_EN.pdf"
EN_PNG = OUT / "02_Standard_Factor_Two_Level_Tree_EN.png"

CN_XLSX = OUT / "03_标准因素频次统计表_中文.xlsx"
EN_XLSX = OUT / "04_Standard_Factor_Frequency_Tables_EN.xlsx"
BASE_XLSX = OUT / "05_可视化统计底表.xlsx"
ISSUE_XLSX = OUT / "数据一致性异常.xlsx"

REQUIRED_COLUMNS = [
    "第一作者",
    "发表年份",
    "原始因素名称",
    "原始因素名称（英文）",
    "一级分类",
    "标准因素名称",
    "标准因素名称（英文）",
    "主要参考标准",
]

CATEGORY_EN = {
    "卫生健康服务与系统": "Health Services and Systems",
    "水、环境卫生与个人卫生（WASH）": "Water, Sanitation and Hygiene (WASH)",
    "食物安全与营养": "Food Security and Nutrition",
    "住所、安置与人口流离失所": "Shelter, Settlement and Population Displacement",
    "保护、冲突与危险暴露": "Protection, Conflict and Hazard Exposure",
    "人道准入、交通物流与关键基础设施": "Humanitarian Access, Transport, Logistics and Critical Infrastructure",
    "社会经济、治理与权利环境": "Socioeconomic, Governance, and Rights Environment",
    "其他/待人工复核": "Other / Pending Manual Review",
}

CATEGORY_COLORS = {
    "卫生健康服务与系统": {"dark": "#24557A", "light": "#DCEAF3"},
    "食物安全与营养": {"dark": "#A45C40", "light": "#F4E4DC"},
    "住所、安置与人口流离失所": {"dark": "#6B568F", "light": "#EAE4F2"},
    "保护、冲突与危险暴露": {"dark": "#963F59", "light": "#F3E1E7"},
    "水、环境卫生与个人卫生（WASH）": {"dark": "#2B776F", "light": "#DCEEEB"},
    "社会经济、治理与权利环境": {"dark": "#6D6A35", "light": "#EEECD8"},
    "人道准入、交通物流与关键基础设施": {"dark": "#4D6878", "light": "#E1E9ED"},
    "其他/待人工复核": {"dark": "#6B7280", "light": "#ECEFF1"},
}

CN_TITLE = "战争期间患病人员死亡增加相关因素的两级树状结构"
CN_SUBTITLE = "仅显示出现频次≥2的标准因素；一级分类及类内标准因素均按频次降序排列"
CN_NOTE = "注：一级分类“总n”为该分类全部有效记录数；“图示n”为频次≥2标准因素对应的记录数。出现频次为1的标准因素未在图中显示，详见中文附表。"
EN_TITLE = "Two-Level Tree of Factors Associated with Increased Mortality Among Ill Persons During War"
EN_SUBTITLE = "Standard factors with frequency ≥2; primary categories and factors are ordered by descending frequency"
EN_NOTE = "Note: “Total n” represents all valid records within the primary category, whereas “Shown n” includes only records assigned to standard factors occurring at least twice. Standard factors occurring once are listed in the English appendix table."

CN_REGULAR_CANDIDATES = [
    Path(r"C:\Windows\Fonts\msyh.ttc"),
    Path(r"C:\Windows\Fonts\simhei.ttf"),
]
CN_BOLD_CANDIDATES = [
    Path(r"C:\Windows\Fonts\msyhbd.ttc"),
    Path(r"C:\Windows\Fonts\simhei.ttf"),
]
EN_REGULAR_CANDIDATES = [Path(r"C:\Windows\Fonts\arial.ttf")]
EN_BOLD_CANDIDATES = [Path(r"C:\Windows\Fonts\arialbd.ttf")]

# Publication-oriented figure settings.  Layout coordinates are typographic
# points (72 points = 1 inch); the raster export scales the same geometry to
# 600 dpi while SVG and PDF remain editable/vector outputs.
FIGURE_WIDTH_PT = 504.0  # 7.0 inches: fits a full-width journal text block.
PNG_DPI = 600
MEASURE_SCALE = 10


def now_iso() -> str:
    return datetime.now().astimezone().isoformat(timespec="seconds")


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest().upper()


def file_state(path: Path) -> dict[str, Any]:
    stat = path.stat()
    return {
        "path": str(path),
        "size": stat.st_size,
        "mtime_ns": stat.st_mtime_ns,
        "sha256": sha256(path),
    }


def pick_font(candidates: Iterable[Path]) -> Path:
    for path in candidates:
        if path.exists():
            return path
    raise FileNotFoundError(f"No usable font found among: {list(candidates)}")


def norm_text(value: Any) -> str:
    if pd.isna(value):
        return ""
    return re.sub(r"\s+", " ", str(value).strip())


def loose_key(value: Any) -> str:
    text = unicodedata.normalize("NFKC", norm_text(value)).casefold()
    return re.sub(r"[^\w\u4e00-\u9fff]+", "", text)


def sorted_unique(values: Iterable[Any]) -> list[str]:
    return sorted({norm_text(v) for v in values if norm_text(v)}, key=lambda x: x.casefold())


def fit_lines(text: str, font: ImageFont.FreeTypeFont, max_width: int, english: bool) -> list[str]:
    text = norm_text(text)
    if not text:
        return [""]
    if english:
        tokens = text.split()
        lines: list[str] = []
        current = ""
        for token in tokens:
            trial = token if not current else f"{current} {token}"
            if font.getlength(trial) <= max_width:
                current = trial
            else:
                if current:
                    lines.append(current)
                if font.getlength(token) <= max_width:
                    current = token
                else:
                    chunk = ""
                    for char in token:
                        if chunk and font.getlength(chunk + char) > max_width:
                            lines.append(chunk)
                            chunk = char
                        else:
                            chunk += char
                    current = chunk
        if current:
            lines.append(current)
        return lines

    tokens = re.findall(r"[A-Za-z0-9][A-Za-z0-9()/_–—+.-]*|\s+|.", text)
    lines = []
    current = ""
    for token in tokens:
        if token.isspace() and not current:
            continue
        trial = current + token
        if font.getlength(trial) <= max_width:
            current = trial
        else:
            if current.strip():
                lines.append(current.strip())
            current = token.strip()
    if current.strip():
        lines.append(current.strip())
    return lines or [text]


def point_font(path: Path, size_pt: float) -> ImageFont.FreeTypeFont:
    """Return a high-resolution Pillow font for point-based layout measuring."""
    return ImageFont.truetype(str(path), max(1, round(size_pt * MEASURE_SCALE)))


def fit_lines_pt(text: str, font: ImageFont.FreeTypeFont, max_width_pt: float, english: bool) -> list[str]:
    return fit_lines(text, font, max(1, round(max_width_pt * MEASURE_SCALE)), english)


def add_issue(
    issues: list[dict[str, Any]],
    issue_type: str,
    severity: str,
    subject: str,
    rows: str,
    detail: str,
    handling: str,
) -> None:
    issues.append(
        {
            "问题类型": issue_type,
            "严重程度": severity,
            "涉及对象": subject,
            "涉及行号": rows,
            "问题详情": detail,
            "处理方式": handling,
        }
    )


def load_and_analyze() -> dict[str, Any]:
    if not INPUT.exists():
        raise FileNotFoundError(f"正式输入文件不存在：{INPUT}")
    source_baseline = file_state(INPUT)
    codebook_baseline = file_state(CODEBOOK) if CODEBOOK.exists() else None

    xls = pd.ExcelFile(INPUT)
    if "因素提取" not in xls.sheet_names:
        raise RuntimeError("正式输入中不存在“因素提取”工作表。")
    df = pd.read_excel(INPUT, sheet_name="因素提取", dtype=object)
    missing_columns = [c for c in REQUIRED_COLUMNS if c not in df.columns]
    if "一级分类" in missing_columns or "标准因素名称" in missing_columns:
        raise RuntimeError(f"缺少正式统计关键字段：{missing_columns}")

    issues: list[dict[str, Any]] = []
    for column in missing_columns:
        add_issue(
            issues,
            "缺少预期字段",
            "提示",
            column,
            "",
            f"结果表未包含预期辅助字段“{column}”。",
            "不修改源文件；相应附表字段留空并在日志中记录。",
        )
        df[column] = None

    # Preserve source row numbers for audit and normalize only in-memory display values.
    df["__源Excel行号"] = range(2, len(df) + 2)
    for column in REQUIRED_COLUMNS:
        if column != "发表年份":
            df[column] = df[column].map(norm_text)

    valid_mask = (df["一级分类"] != "") & (df["标准因素名称"] != "")
    valid = df.loc[valid_mask].copy()
    invalid = df.loc[~valid_mask].copy()
    for _, row in invalid.iterrows():
        missing = []
        if not row["一级分类"]:
            missing.append("一级分类")
        if not row["标准因素名称"]:
            missing.append("标准因素名称")
        add_issue(
            issues,
            "有效记录关键字段缺失",
            "错误",
            row.get("原始因素名称", ""),
            str(row["__源Excel行号"]),
            f"缺失字段：{'、'.join(missing)}。",
            "该行不进入正式频次统计；源文件保持不变。",
        )

    # Canonical codebook lookup is used only for consistency checking and English naming.
    codebook_map: dict[str, dict[str, str]] = {}
    if CODEBOOK.exists():
        cb = pd.read_excel(CODEBOOK, sheet_name="代码本", dtype=object)
        for _, row in cb.iterrows():
            name = norm_text(row.get("标准因素名称"))
            if not name:
                continue
            codebook_map[name] = {
                "category": norm_text(row.get("一级分类")),
                "english": norm_text(row.get("标准因素名称（英文）")),
            }

    # 1–3: cross-category and bilingual one-to-many checks.
    for name, group in valid.groupby("标准因素名称", sort=False):
        cats = sorted_unique(group["一级分类"])
        ens = sorted_unique(group["标准因素名称（英文）"])
        if len(cats) > 1:
            add_issue(
                issues,
                "标准因素跨一级分类",
                "阻断",
                name,
                "、".join(map(str, group["__源Excel行号"].tolist())),
                f"同一标准因素对应多个一级分类：{'；'.join(cats)}。",
                "不得静默处理；需依据正式代码本人工确认。",
            )
        if len(ens) > 1:
            add_issue(
                issues,
                "中文标准因素对应多个英文名称",
                "警告",
                name,
                "、".join(map(str, group["__源Excel行号"].tolist())),
                f"英文名称变体：{'；'.join(ens)}。",
                "图表优先使用正式代码本英文名称；源文件保持不变。",
            )

    for en_name, group in valid[valid["标准因素名称（英文）"] != ""].groupby("标准因素名称（英文）", sort=False):
        names = sorted_unique(group["标准因素名称"])
        if len(names) > 1:
            add_issue(
                issues,
                "英文标准因素对应多个中文名称",
                "警告",
                en_name,
                "、".join(map(str, group["__源Excel行号"].tolist())),
                f"中文名称：{'；'.join(names)}。",
                "仅记录，不修改源文件。",
            )

    # 4: formatting-only collisions.
    for column, label in [("标准因素名称", "中文标准因素"), ("标准因素名称（英文）", "英文标准因素")]:
        buckets: dict[str, set[str]] = defaultdict(set)
        for value in valid[column]:
            if norm_text(value):
                buckets[loose_key(value)].add(norm_text(value))
        for variants in buckets.values():
            if len(variants) > 1:
                add_issue(
                    issues,
                    "仅格式差异的标准因素",
                    "警告",
                    label,
                    "",
                    f"发现可能仅空格、标点或大小写不同的表达：{'；'.join(sorted(variants))}。",
                    "仅记录；不得在本任务中合并或改名。",
                )

    # 7: English name missing.
    for _, row in valid[valid["标准因素名称（英文）"] == ""].iterrows():
        add_issue(
            issues,
            "英文标准因素名称缺失",
            "警告",
            row["标准因素名称"],
            str(row["__源Excel行号"]),
            "正式结果中的标准因素英文名称为空。",
            "优先检查正式代码本；若仍缺失，仅使用项目忠实译名并记录。",
        )

    missing_year = valid[valid["发表年份"].isna() | (valid["发表年份"].map(norm_text) == "")]
    if not missing_year.empty:
        add_issue(
            issues,
            "发表年份缺失",
            "提示",
            "；".join(sorted_unique(missing_year["标准因素名称"])),
            "、".join(map(str, missing_year["__源Excel行号"].tolist())),
            f"共有{len(missing_year)}条有效记录缺少发表年份，其中频次1附表对应行将保留空值。",
            "不推测年份，不修改源文件；在附表和异常清单中保留空值并记录。",
        )

    # Codebook alignment check.  The result workbook remains the statistical authority.
    for name, group in valid.groupby("标准因素名称", sort=False):
        category = sorted_unique(group["一级分类"])[0]
        source_en = sorted_unique(group["标准因素名称（英文）"])
        cb_item = codebook_map.get(name)
        if not cb_item:
            add_issue(
                issues,
                "结果表标准因素未收录于代码本",
                "提示",
                name,
                "、".join(map(str, group["__源Excel行号"].tolist())),
                "该标准因素存在于正式结果表，但未在当前代码本中找到。",
                "按任务边界以正式结果表为准统计；不新增、不改名、不重新分类。",
            )
        else:
            if cb_item["category"] and cb_item["category"] != category:
                add_issue(
                    issues,
                    "结果表与代码本一级分类不一致",
                    "警告",
                    name,
                    "、".join(map(str, group["__源Excel行号"].tolist())),
                    f"结果表：{category}；代码本：{cb_item['category']}。",
                    "仅记录；本任务仍以正式结果表为统计依据。",
                )
            if cb_item["english"] and source_en and cb_item["english"] not in source_en:
                add_issue(
                    issues,
                    "结果表与代码本英文名称不一致",
                    "警告",
                    name,
                    "、".join(map(str, group["__源Excel行号"].tolist())),
                    f"结果表：{'；'.join(source_en)}；代码本：{cb_item['english']}。",
                    "英文图表使用代码本名称；源文件保持不变。",
                )

    # 8: exact duplicate row groups across the eight formal fields.  Keep every row in counts.
    duplicate_details: list[dict[str, Any]] = []
    duplicate_mask = df.duplicated(subset=REQUIRED_COLUMNS, keep=False)
    if duplicate_mask.any():
        dup = df.loc[duplicate_mask].copy()
        group_no = 0
        for _, group in dup.groupby(REQUIRED_COLUMNS, dropna=False, sort=False):
            group_no += 1
            row_numbers = [int(x) for x in group["__源Excel行号"].tolist()]
            first = group.iloc[0]
            detail = {
                "重复组编号": group_no,
                "重复行数": len(group),
                "源Excel行号": "、".join(map(str, row_numbers)),
            }
            for column in REQUIRED_COLUMNS:
                value = first[column]
                if pd.isna(value):
                    value = ""
                elif hasattr(value, "item"):
                    try:
                        value = value.item()
                    except Exception:
                        pass
                detail[column] = value
            duplicate_details.append(detail)
            factor_name = norm_text(first["标准因素名称"])
            factor_total_frequency = int((valid["标准因素名称"] == factor_name).sum())
            threshold_note = ""
            if factor_total_frequency == len(group) == 2:
                threshold_note = " 该因素的正式频次n=2完全由本重复组构成，因此按既定逐行口径进入树状图。"
            add_issue(
                issues,
                "完全重复数据行",
                "提示",
                factor_name,
                detail["源Excel行号"],
                f"8个正式字段完全相同，共{len(group)}行。{threshold_note}",
                "依据用户规定的逐行频次口径全部保留并计数；不修改源文件。",
            )

    blocking = [i for i in issues if i["严重程度"] == "阻断"]
    if blocking:
        WORK.mkdir(parents=True, exist_ok=True)
        DATA_JSON.write_text(json.dumps({"status": "blocked", "issues": issues}, ensure_ascii=False, indent=2), encoding="utf-8")
        raise RuntimeError("发现无法安全解决的跨一级分类冲突，已暂停正式树状图生成。")

    # Canonical English names: codebook first, then most frequent non-empty source value.
    canonical_en: dict[str, str] = {}
    for name, group in valid.groupby("标准因素名称", sort=False):
        cb_en = codebook_map.get(name, {}).get("english", "")
        source_values = [norm_text(x) for x in group["标准因素名称（英文）"] if norm_text(x)]
        if cb_en:
            canonical_en[name] = cb_en
        elif source_values:
            counts = Counter(source_values)
            canonical_en[name] = sorted(counts, key=lambda x: (-counts[x], x.casefold()))[0]
        else:
            canonical_en[name] = name

    factor_rows: list[dict[str, Any]] = []
    for (category, name), group in valid.groupby(["一级分类", "标准因素名称"], sort=False):
        n = int(len(group))
        source_terms = {norm_text(x) for x in group["原始因素名称"] if norm_text(x)}
        factor_rows.append(
            {
                "category_cn": category,
                "category_en": CATEGORY_EN.get(category, category),
                "factor_cn": name,
                "factor_en": canonical_en[name],
                "frequency": n,
                "distinct_source_terms": len(source_terms),
                "source_rows": [int(x) for x in group["__源Excel行号"].tolist()],
            }
        )

    category_totals = Counter(valid["一级分类"])
    categories = sorted(category_totals, key=lambda c: (-category_totals[c], CATEGORY_EN.get(c, c).casefold()))
    category_order = {c: idx + 1 for idx, c in enumerate(categories)}

    factor_rows.sort(
        key=lambda r: (
            category_order[r["category_cn"]],
            -r["frequency"],
            r["factor_en"].casefold(),
            r["factor_cn"],
        )
    )
    within_counter: dict[str, int] = defaultdict(int)
    for row in factor_rows:
        within_counter[row["category_cn"]] += 1
        row["within_category_rank"] = within_counter[row["category_cn"]]
        row["category_order"] = category_order[row["category_cn"]]
        row["shown"] = row["frequency"] >= 2
        row["pct_all"] = row["frequency"] / len(valid) if len(valid) else 0
        row["pct_category"] = row["frequency"] / category_totals[row["category_cn"]]

    overall_sorted = sorted(factor_rows, key=lambda r: (-r["frequency"], r["factor_en"].casefold(), r["factor_cn"]))
    overall_rank = {(r["category_cn"], r["factor_cn"]): i + 1 for i, r in enumerate(overall_sorted)}
    for row in factor_rows:
        row["overall_rank"] = overall_rank[(row["category_cn"], row["factor_cn"])]

    category_summary: list[dict[str, Any]] = []
    for category in categories:
        rows = [r for r in factor_rows if r["category_cn"] == category]
        shown_rows = [r for r in rows if r["shown"]]
        singleton_rows = [r for r in rows if r["frequency"] == 1]
        category_summary.append(
            {
                "order": category_order[category],
                "category_cn": category,
                "category_en": CATEGORY_EN.get(category, category),
                "total_frequency": int(category_totals[category]),
                "shown_frequency": sum(r["frequency"] for r in shown_rows),
                "ge2_factor_count": len(shown_rows),
                "singleton_factor_count": len(singleton_rows),
                "total_factor_count": len(rows),
                "pct_all": category_totals[category] / len(valid) if len(valid) else 0,
                "colors": CATEGORY_COLORS.get(category, CATEGORY_COLORS["其他/待人工复核"]),
            }
        )

    main_rows = [r.copy() for r in factor_rows if r["shown"]]
    singletons = [r.copy() for r in factor_rows if r["frequency"] == 1]

    # Attach the one source record represented by each singleton factor.
    singleton_source: dict[tuple[str, str], dict[str, Any]] = {}
    for (category, name), group in valid.groupby(["一级分类", "标准因素名称"], sort=False):
        if len(group) == 1:
            row = group.iloc[0]
            year = row["发表年份"]
            if pd.isna(year):
                year = ""
            elif isinstance(year, float) and year.is_integer():
                year = int(year)
            singleton_source[(category, name)] = {
                "source_factor_cn": norm_text(row["原始因素名称"]),
                "source_factor_en": norm_text(row["原始因素名称（英文）"]),
                "first_author": norm_text(row["第一作者"]),
                "publication_year": year,
                "reference": norm_text(row["主要参考标准"]),
                "source_excel_row": int(row["__源Excel行号"]),
            }
    for row in singletons:
        row.update(singleton_source[(row["category_cn"], row["factor_cn"])])

    # Presentation-only singleton grouping comes from an auditable complete mapping.
    group_path = WORK / "singleton_group_overrides.json"
    if not group_path.exists():
        group_path = WORK / "agent_singleton_groups.json"
    group_map: dict[tuple[str, str], tuple[str, str]] = {}
    if group_path.exists():
        raw_groups = json.loads(group_path.read_text(encoding="utf-8"))
        if isinstance(raw_groups, list):
            raw_items = raw_groups
        else:
            raw_items = raw_groups.get("groups", raw_groups.get("items", []))
        for item in raw_items:
            group_map[(item["一级分类"], item["标准因素名称"])] = (
                item.get("展示小标题中文", item.get("展示小标题（非正式）", "")),
                item.get("展示小标题英文", item.get("展示小标题（英文，非分类）", "")),
            )

    singleton_counts = Counter(r["category_cn"] for r in singletons)
    for row in singletons:
        key = (row["category_cn"], row["factor_cn"])
        if singleton_counts[row["category_cn"]] <= 5:
            subgroup_cn, subgroup_en = group_map.get(key, ("", ""))
        else:
            subgroup_cn, subgroup_en = group_map.get(key, ("其他具体因素", "Other specific factors"))
        row["subgroup_cn"] = subgroup_cn
        row["subgroup_en"] = subgroup_en

    # Appendix order: category order, subgroup size desc, subgroup English, factor English.
    subgroup_counts = Counter((r["category_cn"], r["subgroup_cn"], r["subgroup_en"]) for r in singletons)
    singletons.sort(
        key=lambda r: (
            category_order[r["category_cn"]],
            -subgroup_counts[(r["category_cn"], r["subgroup_cn"], r["subgroup_en"])],
            r["subgroup_en"].casefold(),
            r["subgroup_cn"],
            r["factor_en"].casefold(),
            r["factor_cn"],
        )
    )

    result = {
        "status": "prepared",
        "generated_at": now_iso(),
        "source": source_baseline,
        "codebook": codebook_baseline,
        "input_sheet": "因素提取",
        "columns": REQUIRED_COLUMNS,
        "metrics": {
            "source_data_rows": len(df),
            "valid_records": len(valid),
            "invalid_records": len(invalid),
            "category_count": len(categories),
            "standard_factor_count": len(factor_rows),
            "ge2_factor_count": len(main_rows),
            "singleton_factor_count": len(singletons),
            "shown_frequency": sum(r["frequency"] for r in main_rows),
            "tree_node_count": len(categories) + len(main_rows),
            "duplicate_row_instances": int(duplicate_mask.sum()),
            "duplicate_extra_rows": int(duplicate_mask.sum() - len(duplicate_details)),
            "duplicate_group_count": len(duplicate_details),
        },
        "category_summary": category_summary,
        "factor_summary": factor_rows,
        "main_rows": main_rows,
        "singleton_rows": singletons,
        "issues": issues,
        "duplicate_details": duplicate_details,
        "palette": CATEGORY_COLORS,
        "titles": {
            "cn_title": CN_TITLE,
            "cn_subtitle": CN_SUBTITLE,
            "cn_note": CN_NOTE,
            "en_title": EN_TITLE,
            "en_subtitle": EN_SUBTITLE,
            "en_note": EN_NOTE,
        },
        "paths": {
            "cn_svg": str(CN_SVG),
            "cn_pdf": str(CN_PDF),
            "cn_png": str(CN_PNG),
            "en_svg": str(EN_SVG),
            "en_pdf": str(EN_PDF),
            "en_png": str(EN_PNG),
            "cn_xlsx": str(CN_XLSX),
            "en_xlsx": str(EN_XLSX),
            "base_xlsx": str(BASE_XLSX),
            "issue_xlsx": str(ISSUE_XLSX),
            "log": str(LOG_PATH),
        },
    }
    return result


def build_layout(data: dict[str, Any]) -> dict[str, Any]:
    width = FIGURE_WIDTH_PT
    margin_x = 11.0
    content_w = width - 2 * margin_x
    panel_gap = 6.0
    row_gap = 5.0
    node_gap = 1.8
    panel_pad_x = 3.5
    tree_gutter = 4.0
    connector_gap = 7.0

    sizes = {
        "title_cn": 10.4,
        "title_en": 9.2,
        "subtitle": 6.1,
        "legend": 5.8,
        "category_cn": 6.6,
        "category_en": 6.2,
        "category_line": 7.2,
        "factor_cn": 6.0,
        "factor_en": 5.8,
        "factor_line": 6.7,
        "frequency": 5.4,
        "note": 5.4,
    }
    font_cn_regular = point_font(pick_font(CN_REGULAR_CANDIDATES), sizes["factor_cn"])
    font_en_regular = point_font(pick_font(EN_REGULAR_CANDIDATES), sizes["factor_en"])
    font_cn_bold = point_font(pick_font(CN_BOLD_CANDIDATES), sizes["category_cn"])
    font_en_bold = point_font(pick_font(EN_BOLD_CANDIDATES), sizes["category_en"])
    title_font_cn = point_font(pick_font(CN_BOLD_CANDIDATES), sizes["title_cn"])
    title_font_en = point_font(pick_font(EN_BOLD_CANDIDATES), sizes["title_en"])
    subtitle_font_cn = point_font(pick_font(CN_REGULAR_CANDIDATES), sizes["subtitle"])
    subtitle_font_en = point_font(pick_font(EN_REGULAR_CANDIDATES), sizes["subtitle"])
    note_font_cn = point_font(pick_font(CN_REGULAR_CANDIDATES), sizes["note"])
    note_font_en = point_font(pick_font(EN_REGULAR_CANDIDATES), sizes["note"])

    cn_title_lines = fit_lines_pt(CN_TITLE, title_font_cn, content_w, False)
    en_title_lines = fit_lines_pt(EN_TITLE, title_font_en, content_w, True)
    title_line_h = 11.6
    title_h = max(len(cn_title_lines), len(en_title_lines)) * title_line_h
    title_y = 6.0 + title_h / 2

    cn_subtitle_lines = fit_lines_pt(CN_SUBTITLE, subtitle_font_cn, content_w, False)
    en_subtitle_lines = fit_lines_pt(EN_SUBTITLE, subtitle_font_en, content_w, True)
    subtitle_line_h = 7.0
    subtitle_h = max(len(cn_subtitle_lines), len(en_subtitle_lines)) * subtitle_line_h
    subtitle_y = 6.0 + title_h + 1.5 + subtitle_h / 2
    legend_y = 6.0 + title_h + 1.5 + subtitle_h + 6.0
    content_top = legend_y + 8.0

    main_by_category: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in data["main_rows"]:
        main_by_category[row["category_cn"]].append(row)

    categories = data["category_summary"]
    # The current seven-category dataset is laid out in four compact reading
    # bands.  The sequence remains the formal frequency order.  Wide bands use
    # multiple newspaper-style factor columns; smaller categories share a row.
    if len(categories) == 7:
        row_specs = [
            [0],
            [1, 2],
            [3],
            [4, 5, 6],
        ]
        factor_column_specs = {
            0: 4,
            1: 1,
            2: 1,
            3: 4,
            4: 1,
            5: 1,
            6: 1,
        }
    else:
        row_specs = [[0]] if categories else []
        row_specs.extend([list(range(i, min(i + 3, len(categories)))) for i in range(1, len(categories), 3)])
        factor_column_specs = {
            i: (min(5, max(1, math.ceil(len(main_by_category[c["category_cn"]]) / 12))) if i == 0 else (2 if len(main_by_category[c["category_cn"]]) > 10 else 1))
            for i, c in enumerate(categories)
        }

    category_nodes: list[dict[str, Any]] = []
    factor_nodes: list[dict[str, Any]] = []
    panel_rows: list[dict[str, Any]] = []
    factor_index = 0
    y = content_top

    for row_no, indices in enumerate(row_specs, start=1):
        columns_in_row = len(indices)
        panel_w = (content_w - panel_gap * (columns_in_row - 1)) / columns_in_row
        row_panels: list[dict[str, Any]] = []
        row_height = 0.0
        for position, category_index in enumerate(indices):
            category = categories[category_index]
            cat_cn = category["category_cn"]
            factors = main_by_category[cat_cn]
            factor_cols = min(max(1, factor_column_specs.get(category_index, 1)), max(1, len(factors)))
            x = margin_x + position * (panel_w + panel_gap)

            count_cn = f"总n={category['total_frequency']} · 图示n={category['shown_frequency']}"
            count_en = f"Total n={category['total_frequency']} · Shown n={category['shown_frequency']}"
            cn_cat_lines = fit_lines_pt(cat_cn, font_cn_bold, panel_w - 9.0, False) + [count_cn]
            en_cat_lines = fit_lines_pt(category["category_en"], font_en_bold, panel_w - 9.0, True) + [count_en]
            category_h = max(17.0, max(len(cn_cat_lines), len(en_cat_lines)) * sizes["category_line"] + 4.0)
            bus_y = y + category_h + 3.2
            body_y = y + category_h + connector_gap

            inner_w = panel_w - 2 * panel_pad_x
            col_gap = 3.4
            col_w = (inner_w - col_gap * (factor_cols - 1)) / factor_cols
            node_w = col_w - tree_gutter
            frequency_w = 18.0
            factor_pad = 2.3
            name_w = max(18.0, node_w - 2 * factor_pad - frequency_w)
            rows_per_col = max(1, math.ceil(len(factors) / factor_cols))
            col_y = [body_y for _ in range(factor_cols)]
            nodes_by_col: list[list[dict[str, Any]]] = [[] for _ in range(factor_cols)]

            for i, factor in enumerate(factors):
                factor_index += 1
                col = min(i // rows_per_col, factor_cols - 1)
                col_left = x + panel_pad_x + col * (col_w + col_gap)
                node_x = col_left + tree_gutter
                cn_lines = fit_lines_pt(factor["factor_cn"], font_cn_regular, name_w, False)
                en_lines = fit_lines_pt(factor["factor_en"], font_en_regular, name_w, True)
                node_h = max(11.8, max(len(cn_lines), len(en_lines)) * sizes["factor_line"] + 3.2)
                node = {
                    "node_id": f"f-{factor_index:03d}",
                    "category_cn": cat_cn,
                    "factor_cn": factor["factor_cn"],
                    "factor_en": factor["factor_en"],
                    "frequency": factor["frequency"],
                    "rank": factor["within_category_rank"],
                    "x": node_x,
                    "y": col_y[col],
                    "w": node_w,
                    "h": node_h,
                    "name_x": node_x + factor_pad,
                    "name_w": name_w,
                    "frequency_x": node_x + node_w - factor_pad,
                    "cn_lines": cn_lines,
                    "en_lines": en_lines,
                    "colors": category["colors"],
                }
                factor_nodes.append(node)
                nodes_by_col[col].append(node)
                col_y[col] += node_h + node_gap

            branches: list[dict[str, float]] = []
            for col, nodes in enumerate(nodes_by_col):
                if not nodes:
                    continue
                col_left = x + panel_pad_x + col * (col_w + col_gap)
                branches.append(
                    {
                        "trunk_x": col_left + 1.6,
                        "first_child_y": nodes[0]["y"] + nodes[0]["h"] / 2,
                        "last_child_y": nodes[-1]["y"] + nodes[-1]["h"] / 2,
                        "node_ids": [node["node_id"] for node in nodes],
                    }
                )

            body_bottom = max((v - node_gap for v in col_y), default=body_y)
            panel_h = body_bottom - y + 2.5
            category_node = {
                "node_id": f"c-{category['order']}",
                "category_cn": cat_cn,
                "category_en": category["category_en"],
                "x": x,
                "y": y,
                "w": panel_w,
                "h": category_h,
                "header_center_x": x + panel_w / 2,
                "header_bottom_y": y + category_h,
                "bus_y": bus_y,
                "branches": branches,
                "cn_lines": cn_cat_lines,
                "en_lines": en_cat_lines,
                "colors": category["colors"],
                "panel_h": panel_h,
                "factor_columns": factor_cols,
            }
            category_nodes.append(category_node)
            row_panels.append(category_node)
            row_height = max(row_height, panel_h)

        panel_rows.append({"row": row_no, "y": y, "height": row_height, "categories": [p["category_cn"] for p in row_panels]})
        y += row_height + row_gap

    cn_note_lines = fit_lines_pt(CN_NOTE, note_font_cn, content_w, False)
    en_note_lines = fit_lines_pt(EN_NOTE, note_font_en, content_w, True)
    note_line_h = 6.4
    note_h = max(len(cn_note_lines), len(en_note_lines)) * note_line_h
    note_y = y + 2.5 + note_h / 2
    height = note_y + note_h / 2 + 6.0

    return {
        "design": "compact-paper-multipanel",
        "dpi": PNG_DPI,
        "width": round(width, 2),
        "height": round(height, 2),
        "physical_width_in": round(width / 72.0, 3),
        "physical_height_in": round(height / 72.0, 3),
        "sizes": sizes,
        "category_nodes": category_nodes,
        "factor_nodes": factor_nodes,
        "panel_rows": panel_rows,
        "cn_title_lines": cn_title_lines,
        "en_title_lines": en_title_lines,
        "title_y": title_y,
        "cn_subtitle_lines": cn_subtitle_lines,
        "en_subtitle_lines": en_subtitle_lines,
        "subtitle_y": subtitle_y,
        "legend_y": legend_y,
        "cn_note_lines": cn_note_lines,
        "en_note_lines": en_note_lines,
        "note_y": note_y,
    }


def _text_svg(
    x: float,
    y: float,
    lines: list[str],
    size: float,
    color: str,
    family: str,
    weight: str = "400",
    anchor: str = "middle",
    line_height: float | None = None,
) -> str:
    line_height = line_height or size * 1.2
    start_y = y - (len(lines) - 1) * line_height / 2
    tspans = "".join(
        f'<tspan x="{x:.1f}" y="{start_y + i * line_height:.1f}">{html.escape(line)}</tspan>'
        for i, line in enumerate(lines)
    )
    return (
        f'<text x="{x:.1f}" y="{y:.1f}" text-anchor="{anchor}" dominant-baseline="middle" '
        f'font-family="{family}" font-size="{size}" font-weight="{weight}" fill="{color}">{tspans}</text>'
    )


def render_svg(data: dict[str, Any], layout: dict[str, Any], locale: str, path: Path) -> None:
    width, height = layout["width"], layout["height"]
    family = "Microsoft YaHei, SimHei, Noto Sans CJK SC, sans-serif" if locale == "cn" else "Arial, Helvetica, sans-serif"
    title = CN_TITLE if locale == "cn" else EN_TITLE
    subtitle = CN_SUBTITLE if locale == "cn" else EN_SUBTITLE
    parts = [
        '<?xml version="1.0" encoding="UTF-8"?>',
        f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" viewBox="0 0 {width} {height}" role="img" aria-labelledby="title desc">',
        f'<title id="title">{html.escape(title)}</title>',
        f'<desc id="desc">Two-level node-link tree with {len(layout["category_nodes"])} primary categories and {len(layout["factor_nodes"])} standard factors.</desc>',
        f'<metadata>{{"category_nodes":{len(layout["category_nodes"])},"factor_nodes":{len(layout["factor_nodes"])},"total_nodes":{data["metrics"]["tree_node_count"]}}}</metadata>',
        f'<rect x="0" y="0" width="{width}" height="{height}" fill="#FFFFFF"/>',
        _text_svg(width / 2, 54, [title], 40 if locale == "cn" else 34, "#1F2933", family, "700"),
        _text_svg(width / 2, 108, [subtitle], 24, "#4B5563", family, "400"),
        '<g class="legend">',
        '<rect x="860" y="145" width="44" height="28" rx="6" fill="#24557A"/>',
        _text_svg(1025, 159, ["一级分类" if locale == "cn" else "Primary category"], 21, "#374151", family),
        '<rect x="1240" y="145" width="44" height="28" rx="6" fill="#DCEAF3" stroke="#24557A" stroke-width="2"/>',
        _text_svg(1435, 159, ["标准因素" if locale == "cn" else "Standard factor"], 21, "#374151", family),
        '</g>',
    ]
    factors_by_category: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for node in layout["factor_nodes"]:
        factors_by_category[node["category_cn"]].append(node)
    for cat in layout["category_nodes"]:
        children = factors_by_category[cat["category_cn"]]
        if children:
            dark = cat["colors"]["dark"]
            cat_right = cat["x"] + cat["w"]
            cy = cat["center_y"]
            parts.append(f'<g class="branch" data-node-id="{cat["node_id"]}">')
            parts.append(f'<line x1="{cat_right}" y1="{cy:.1f}" x2="{layout["trunk_x"]}" y2="{cy:.1f}" stroke="{dark}" stroke-width="3"/>')
            parts.append(f'<line x1="{layout["trunk_x"]}" y1="{cat["first_child_y"]:.1f}" x2="{layout["trunk_x"]}" y2="{cat["last_child_y"]:.1f}" stroke="{dark}" stroke-width="3"/>')
            for child in children:
                fy = child["y"] + child["h"] / 2
                parts.append(f'<line x1="{layout["trunk_x"]}" y1="{fy:.1f}" x2="{child["x"]}" y2="{fy:.1f}" stroke="{dark}" stroke-width="2.5"/>')
            parts.append('</g>')
    for cat in layout["category_nodes"]:
        lines = cat["cn_lines"] if locale == "cn" else cat["en_lines"]
        parts.extend(
            [
                f'<g class="category-node" data-node-id="{cat["node_id"]}" data-x="{cat["x"]}" data-y="{cat["y"]:.1f}" data-w="{cat["w"]}" data-h="{cat["h"]}">',
                f'<rect x="{cat["x"]}" y="{cat["y"]:.1f}" width="{cat["w"]}" height="{cat["h"]}" rx="14" fill="{cat["colors"]["dark"]}"/>',
                _text_svg(cat["x"] + cat["w"] / 2, cat["y"] + cat["h"] / 2, lines, 27 if locale == "cn" else 25, "#FFFFFF", family, "700"),
                '</g>',
            ]
        )
    for node in layout["factor_nodes"]:
        lines = node["cn_lines"] if locale == "cn" else node["en_lines"]
        parts.extend(
            [
                f'<g class="factor-node" data-node-id="{node["node_id"]}" data-frequency="{node["frequency"]}" data-x="{node["x"]}" data-y="{node["y"]}" data-w="{node["w"]}" data-h="{node["h"]}">',
                f'<rect x="{node["x"]}" y="{node["y"]}" width="{node["w"]}" height="{node["h"]}" rx="12" fill="{node["colors"]["light"]}" stroke="{node["colors"]["dark"]}" stroke-width="2"/>',
                _text_svg(node["x"] + node["w"] / 2, node["y"] + node["h"] / 2, lines, 26 if locale == "cn" else 25, "#1F2933", family, "400"),
                '</g>',
            ]
        )
    note_lines = layout["cn_note_lines"] if locale == "cn" else layout["en_note_lines"]
    parts.append(_text_svg(width / 2, layout["note_y"], note_lines, 22, "#4B5563", family, "400"))
    parts.append('</svg>')
    path.write_text("\n".join(parts), encoding="utf-8")


def register_pdf_fonts() -> dict[str, str]:
    specs = {
        "cn_regular": ("TreeCN", pick_font(CN_REGULAR_CANDIDATES)),
        "cn_bold": ("TreeCNBold", pick_font(CN_BOLD_CANDIDATES)),
        "en_regular": ("TreeEN", pick_font(EN_REGULAR_CANDIDATES)),
        "en_bold": ("TreeENBold", pick_font(EN_BOLD_CANDIDATES)),
    }
    registered: dict[str, str] = {}
    for key, (name, path) in specs.items():
        try:
            pdfmetrics.registerFont(TTFont(name, str(path), subfontIndex=0))
        except TypeError:
            pdfmetrics.registerFont(TTFont(name, str(path)))
        registered[key] = name
    return registered


def render_pdf(data: dict[str, Any], layout: dict[str, Any], locale: str, path: Path) -> None:
    fonts = register_pdf_fonts()
    scale = 0.75
    width_pt, height_pt = layout["width"] * scale, layout["height"] * scale
    c = canvas.Canvas(str(path), pagesize=(width_pt, height_pt), pageCompression=1)
    c.setTitle(CN_TITLE if locale == "cn" else EN_TITLE)

    def y_pdf(y_top: float) -> float:
        return height_pt - y_top * scale

    def rect(x: float, y: float, w: float, h: float, fill: str, stroke: str | None = None, radius: float = 12) -> None:
        c.setFillColor(HexColor(fill))
        if stroke:
            c.setStrokeColor(HexColor(stroke))
            c.setLineWidth(2 * scale)
            do_stroke = 1
        else:
            do_stroke = 0
        c.roundRect(x * scale, height_pt - (y + h) * scale, w * scale, h * scale, radius * scale, stroke=do_stroke, fill=1)

    def line(x1: float, y1: float, x2: float, y2: float, color: str, width: float) -> None:
        c.setStrokeColor(HexColor(color))
        c.setLineWidth(width * scale)
        c.line(x1 * scale, y_pdf(y1), x2 * scale, y_pdf(y2))

    def text_lines(x: float, y: float, lines: list[str], size: float, color: str, font_name: str) -> None:
        c.setFillColor(HexColor(color))
        c.setFont(font_name, size * scale)
        line_height = size * 1.25
        start = y - (len(lines) - 1) * line_height / 2
        for i, value in enumerate(lines):
            c.drawCentredString(x * scale, y_pdf(start + i * line_height + size * 0.33), value)

    regular = fonts["cn_regular" if locale == "cn" else "en_regular"]
    bold = fonts["cn_bold" if locale == "cn" else "en_bold"]
    title = CN_TITLE if locale == "cn" else EN_TITLE
    subtitle = CN_SUBTITLE if locale == "cn" else EN_SUBTITLE
    c.setFillColor(HexColor("#FFFFFF"))
    c.rect(0, 0, width_pt, height_pt, fill=1, stroke=0)
    text_lines(layout["width"] / 2, 54, [title], 40 if locale == "cn" else 34, "#1F2933", bold)
    text_lines(layout["width"] / 2, 108, [subtitle], 24, "#4B5563", regular)
    rect(860, 145, 44, 28, "#24557A", None, 6)
    text_lines(1025, 159, ["一级分类" if locale == "cn" else "Primary category"], 21, "#374151", regular)
    rect(1240, 145, 44, 28, "#DCEAF3", "#24557A", 6)
    text_lines(1435, 159, ["标准因素" if locale == "cn" else "Standard factor"], 21, "#374151", regular)

    factors_by_category: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for node in layout["factor_nodes"]:
        factors_by_category[node["category_cn"]].append(node)
    for cat in layout["category_nodes"]:
        children = factors_by_category[cat["category_cn"]]
        if children:
            dark = cat["colors"]["dark"]
            line(cat["x"] + cat["w"], cat["center_y"], layout["trunk_x"], cat["center_y"], dark, 3)
            line(layout["trunk_x"], cat["first_child_y"], layout["trunk_x"], cat["last_child_y"], dark, 3)
            for child in children:
                fy = child["y"] + child["h"] / 2
                line(layout["trunk_x"], fy, child["x"], fy, dark, 2.5)
    for cat in layout["category_nodes"]:
        rect(cat["x"], cat["y"], cat["w"], cat["h"], cat["colors"]["dark"], None, 14)
        lines = cat["cn_lines"] if locale == "cn" else cat["en_lines"]
        text_lines(cat["x"] + cat["w"] / 2, cat["y"] + cat["h"] / 2, lines, 27 if locale == "cn" else 25, "#FFFFFF", bold)
    for node in layout["factor_nodes"]:
        rect(node["x"], node["y"], node["w"], node["h"], node["colors"]["light"], node["colors"]["dark"], 12)
        lines = node["cn_lines"] if locale == "cn" else node["en_lines"]
        text_lines(node["x"] + node["w"] / 2, node["y"] + node["h"] / 2, lines, 26 if locale == "cn" else 25, "#1F2933", regular)
    note_lines = layout["cn_note_lines"] if locale == "cn" else layout["en_note_lines"]
    text_lines(layout["width"] / 2, layout["note_y"], note_lines, 22, "#4B5563", regular)
    c.showPage()
    c.save()


def render_png(data: dict[str, Any], layout: dict[str, Any], locale: str, path: Path) -> None:
    image = Image.new("RGB", (layout["width"], layout["height"]), "#FFFFFF")
    draw = ImageDraw.Draw(image)
    regular_path = pick_font(CN_REGULAR_CANDIDATES if locale == "cn" else EN_REGULAR_CANDIDATES)
    bold_path = pick_font(CN_BOLD_CANDIDATES if locale == "cn" else EN_BOLD_CANDIDATES)
    regular = ImageFont.truetype(str(regular_path), 26 if locale == "cn" else 25)
    bold = ImageFont.truetype(str(bold_path), 27 if locale == "cn" else 25)
    title_font = ImageFont.truetype(str(bold_path), 40 if locale == "cn" else 34)
    subtitle_font = ImageFont.truetype(str(regular_path), 24)
    legend_font = ImageFont.truetype(str(regular_path), 21)
    note_font = ImageFont.truetype(str(regular_path), 22)

    def centered(lines: list[str], x: float, y: float, font: ImageFont.FreeTypeFont, fill: str, line_height: int | None = None) -> None:
        lh = line_height or int(font.size * 1.25)
        total = len(lines) * lh
        top = y - total / 2
        for i, value in enumerate(lines):
            box = draw.textbbox((0, 0), value, font=font)
            tw = box[2] - box[0]
            th = box[3] - box[1]
            draw.text((x - tw / 2, top + i * lh + (lh - th) / 2 - box[1]), value, font=font, fill=fill)

    title = CN_TITLE if locale == "cn" else EN_TITLE
    subtitle = CN_SUBTITLE if locale == "cn" else EN_SUBTITLE
    centered([title], layout["width"] / 2, 54, title_font, "#1F2933")
    centered([subtitle], layout["width"] / 2, 108, subtitle_font, "#4B5563")
    draw.rounded_rectangle((860, 145, 904, 173), radius=6, fill="#24557A")
    centered(["一级分类" if locale == "cn" else "Primary category"], 1025, 159, legend_font, "#374151")
    draw.rounded_rectangle((1240, 145, 1284, 173), radius=6, fill="#DCEAF3", outline="#24557A", width=2)
    centered(["标准因素" if locale == "cn" else "Standard factor"], 1435, 159, legend_font, "#374151")

    factors_by_category: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for node in layout["factor_nodes"]:
        factors_by_category[node["category_cn"]].append(node)
    for cat in layout["category_nodes"]:
        children = factors_by_category[cat["category_cn"]]
        if children:
            dark = cat["colors"]["dark"]
            draw.line((cat["x"] + cat["w"], cat["center_y"], layout["trunk_x"], cat["center_y"]), fill=dark, width=3)
            draw.line((layout["trunk_x"], cat["first_child_y"], layout["trunk_x"], cat["last_child_y"]), fill=dark, width=3)
            for child in children:
                fy = child["y"] + child["h"] / 2
                draw.line((layout["trunk_x"], fy, child["x"], fy), fill=dark, width=3)
    for cat in layout["category_nodes"]:
        draw.rounded_rectangle((cat["x"], cat["y"], cat["x"] + cat["w"], cat["y"] + cat["h"]), radius=14, fill=cat["colors"]["dark"])
        lines = cat["cn_lines"] if locale == "cn" else cat["en_lines"]
        centered(lines, cat["x"] + cat["w"] / 2, cat["y"] + cat["h"] / 2, bold, "#FFFFFF", 34)
    for node in layout["factor_nodes"]:
        draw.rounded_rectangle(
            (node["x"], node["y"], node["x"] + node["w"], node["y"] + node["h"]),
            radius=12,
            fill=node["colors"]["light"],
            outline=node["colors"]["dark"],
            width=2,
        )
        lines = node["cn_lines"] if locale == "cn" else node["en_lines"]
        centered(lines, node["x"] + node["w"] / 2, node["y"] + node["h"] / 2, regular, "#1F2933", 32)
    note_lines = layout["cn_note_lines"] if locale == "cn" else layout["en_note_lines"]
    centered(note_lines, layout["width"] / 2, layout["note_y"], note_font, "#4B5563", 30)
    image.save(path, format="PNG", dpi=(300, 300), optimize=True)


def render_svg_compact(data: dict[str, Any], layout: dict[str, Any], locale: str, path: Path) -> None:
    """Render the publication-oriented shared bilingual layout as editable SVG."""
    width, height = layout["width"], layout["height"]
    sizes = layout["sizes"]
    family = "Microsoft YaHei, SimHei, Noto Sans CJK SC, sans-serif" if locale == "cn" else "Arial, Helvetica, sans-serif"
    title = CN_TITLE if locale == "cn" else EN_TITLE
    title_lines = layout["cn_title_lines"] if locale == "cn" else layout["en_title_lines"]
    subtitle_lines = layout["cn_subtitle_lines"] if locale == "cn" else layout["en_subtitle_lines"]
    title_size = sizes["title_cn"] if locale == "cn" else sizes["title_en"]
    category_size = sizes["category_cn"] if locale == "cn" else sizes["category_en"]
    factor_size = sizes["factor_cn"] if locale == "cn" else sizes["factor_en"]
    parts = [
        '<?xml version="1.0" encoding="UTF-8"?>',
        f'<svg xmlns="http://www.w3.org/2000/svg" width="{layout["physical_width_in"]}in" height="{layout["physical_height_in"]}in" viewBox="0 0 {width} {height}" role="img" aria-labelledby="title desc">',
        f'<title id="title">{html.escape(title)}</title>',
        f'<desc id="desc">Compact publication layout: two-level node-link tree with {len(layout["category_nodes"])} primary categories and {len(layout["factor_nodes"])} standard factors.</desc>',
        f'<metadata>{{"design":"{layout["design"]}","png_dpi":{layout["dpi"]},"category_nodes":{len(layout["category_nodes"])},"factor_nodes":{len(layout["factor_nodes"])},"total_nodes":{data["metrics"]["tree_node_count"]}}}</metadata>',
        f'<rect x="0" y="0" width="{width}" height="{height}" fill="#FFFFFF"/>',
        _text_svg(width / 2, layout["title_y"], title_lines, title_size, "#1F2933", family, "700", line_height=11.6),
        _text_svg(width / 2, layout["subtitle_y"], subtitle_lines, sizes["subtitle"], "#4B5563", family, line_height=7.0),
        '<g class="legend">',
        f'<rect x="{width / 2 - 92:.1f}" y="{layout["legend_y"] - 3.0:.1f}" width="9" height="6" rx="1.2" fill="#24557A"/>',
        _text_svg(width / 2 - 78, layout["legend_y"], ["一级分类" if locale == "cn" else "Primary category"], sizes["legend"], "#374151", family, anchor="start"),
        f'<rect x="{width / 2 + 16:.1f}" y="{layout["legend_y"] - 3.0:.1f}" width="9" height="6" rx="1.2" fill="#DCEAF3" stroke="#24557A" stroke-width="0.45"/>',
        _text_svg(width / 2 + 30, layout["legend_y"], ["标准因素" if locale == "cn" else "Standard factor"], sizes["legend"], "#374151", family, anchor="start"),
        '</g>',
    ]
    node_by_id = {node["node_id"]: node for node in layout["factor_nodes"]}
    for cat in layout["category_nodes"]:
        if not cat["branches"]:
            continue
        dark = cat["colors"]["dark"]
        trunk_xs = [branch["trunk_x"] for branch in cat["branches"]]
        bus_left = min(trunk_xs + [cat["header_center_x"]])
        bus_right = max(trunk_xs + [cat["header_center_x"]])
        parts.append(f'<g class="branch" data-node-id="{cat["node_id"]}">')
        parts.append(f'<line x1="{cat["header_center_x"]:.2f}" y1="{cat["header_bottom_y"]:.2f}" x2="{cat["header_center_x"]:.2f}" y2="{cat["bus_y"]:.2f}" stroke="{dark}" stroke-width="0.65"/>')
        parts.append(f'<line x1="{bus_left:.2f}" y1="{cat["bus_y"]:.2f}" x2="{bus_right:.2f}" y2="{cat["bus_y"]:.2f}" stroke="{dark}" stroke-width="0.65"/>')
        for branch in cat["branches"]:
            parts.append(f'<line x1="{branch["trunk_x"]:.2f}" y1="{cat["bus_y"]:.2f}" x2="{branch["trunk_x"]:.2f}" y2="{branch["last_child_y"]:.2f}" stroke="{dark}" stroke-width="0.55"/>')
            for node_id in branch["node_ids"]:
                child = node_by_id[node_id]
                fy = child["y"] + child["h"] / 2
                parts.append(f'<line x1="{branch["trunk_x"]:.2f}" y1="{fy:.2f}" x2="{child["x"]:.2f}" y2="{fy:.2f}" stroke="{dark}" stroke-width="0.45"/>')
        parts.append('</g>')
    for cat in layout["category_nodes"]:
        lines = cat["cn_lines"] if locale == "cn" else cat["en_lines"]
        parts.extend(
            [
                f'<g class="category-node" data-node-id="{cat["node_id"]}" data-x="{cat["x"]:.2f}" data-y="{cat["y"]:.2f}" data-w="{cat["w"]:.2f}" data-h="{cat["h"]:.2f}">',
                f'<rect x="{cat["x"]:.2f}" y="{cat["y"]:.2f}" width="{cat["w"]:.2f}" height="{cat["h"]:.2f}" rx="3.2" fill="{cat["colors"]["dark"]}"/>',
                _text_svg(cat["x"] + cat["w"] / 2, cat["y"] + cat["h"] / 2, lines, category_size, "#FFFFFF", family, "700", line_height=sizes["category_line"]),
                '</g>',
            ]
        )
    for node in layout["factor_nodes"]:
        lines = node["cn_lines"] if locale == "cn" else node["en_lines"]
        parts.extend(
            [
                f'<g class="factor-node" data-node-id="{node["node_id"]}" data-parent="{html.escape(node["category_cn"])}" data-rank="{node["rank"]}" data-frequency="{node["frequency"]}" data-x="{node["x"]:.2f}" data-y="{node["y"]:.2f}" data-w="{node["w"]:.2f}" data-h="{node["h"]:.2f}">',
                f'<rect x="{node["x"]:.2f}" y="{node["y"]:.2f}" width="{node["w"]:.2f}" height="{node["h"]:.2f}" rx="2.2" fill="{node["colors"]["light"]}" stroke="{node["colors"]["dark"]}" stroke-width="0.45"/>',
                _text_svg(node["name_x"], node["y"] + node["h"] / 2, lines, factor_size, "#1F2933", family, anchor="start", line_height=sizes["factor_line"]),
                _text_svg(node["frequency_x"], node["y"] + node["h"] / 2, [f'n={node["frequency"]}'], sizes["frequency"], node["colors"]["dark"], family, "700", anchor="end"),
                '</g>',
            ]
        )
    note_lines = layout["cn_note_lines"] if locale == "cn" else layout["en_note_lines"]
    parts.append(_text_svg(width / 2, layout["note_y"], note_lines, sizes["note"], "#4B5563", family, line_height=6.4))
    parts.append('</svg>')
    path.write_text("\n".join(parts), encoding="utf-8")


def render_pdf_compact(data: dict[str, Any], layout: dict[str, Any], locale: str, path: Path) -> None:
    """Render the same compact point geometry as a single-page vector PDF."""
    fonts = register_pdf_fonts()
    width, height = layout["width"], layout["height"]
    sizes = layout["sizes"]
    c = canvas.Canvas(str(path), pagesize=(width, height), pageCompression=1)
    c.setTitle(CN_TITLE if locale == "cn" else EN_TITLE)

    def y_pdf(y_top: float) -> float:
        return height - y_top

    def rect(x: float, y: float, w: float, h: float, fill: str, stroke: str | None = None, radius: float = 2.2, line_width: float = 0.45) -> None:
        c.setFillColor(HexColor(fill))
        if stroke:
            c.setStrokeColor(HexColor(stroke))
            c.setLineWidth(line_width)
            do_stroke = 1
        else:
            do_stroke = 0
        c.roundRect(x, height - (y + h), w, h, radius, stroke=do_stroke, fill=1)

    def line(x1: float, y1: float, x2: float, y2: float, color: str, width_pt: float) -> None:
        c.setStrokeColor(HexColor(color))
        c.setLineWidth(width_pt)
        c.line(x1, y_pdf(y1), x2, y_pdf(y2))

    def text_lines(x: float, y: float, lines: list[str], size: float, color: str, font_name: str, align: str = "center", line_height: float | None = None) -> None:
        c.setFillColor(HexColor(color))
        c.setFont(font_name, size)
        lh = line_height or size * 1.2
        start = y - (len(lines) - 1) * lh / 2
        for i, value in enumerate(lines):
            baseline = y_pdf(start + i * lh + size * 0.33)
            if align == "left":
                c.drawString(x, baseline, value)
            elif align == "right":
                c.drawRightString(x, baseline, value)
            else:
                c.drawCentredString(x, baseline, value)

    regular = fonts["cn_regular" if locale == "cn" else "en_regular"]
    bold = fonts["cn_bold" if locale == "cn" else "en_bold"]
    title_lines = layout["cn_title_lines"] if locale == "cn" else layout["en_title_lines"]
    subtitle_lines = layout["cn_subtitle_lines"] if locale == "cn" else layout["en_subtitle_lines"]
    title_size = sizes["title_cn"] if locale == "cn" else sizes["title_en"]
    category_size = sizes["category_cn"] if locale == "cn" else sizes["category_en"]
    factor_size = sizes["factor_cn"] if locale == "cn" else sizes["factor_en"]

    c.setFillColor(HexColor("#FFFFFF"))
    c.rect(0, 0, width, height, fill=1, stroke=0)
    text_lines(width / 2, layout["title_y"], title_lines, title_size, "#1F2933", bold, line_height=11.6)
    text_lines(width / 2, layout["subtitle_y"], subtitle_lines, sizes["subtitle"], "#4B5563", regular, line_height=7.0)
    rect(width / 2 - 92, layout["legend_y"] - 3.0, 9, 6, "#24557A", radius=1.2)
    text_lines(width / 2 - 78, layout["legend_y"], ["一级分类" if locale == "cn" else "Primary category"], sizes["legend"], "#374151", regular, "left")
    rect(width / 2 + 16, layout["legend_y"] - 3.0, 9, 6, "#DCEAF3", "#24557A", 1.2)
    text_lines(width / 2 + 30, layout["legend_y"], ["标准因素" if locale == "cn" else "Standard factor"], sizes["legend"], "#374151", regular, "left")

    node_by_id = {node["node_id"]: node for node in layout["factor_nodes"]}
    for cat in layout["category_nodes"]:
        if not cat["branches"]:
            continue
        dark = cat["colors"]["dark"]
        trunk_xs = [branch["trunk_x"] for branch in cat["branches"]]
        bus_left = min(trunk_xs + [cat["header_center_x"]])
        bus_right = max(trunk_xs + [cat["header_center_x"]])
        line(cat["header_center_x"], cat["header_bottom_y"], cat["header_center_x"], cat["bus_y"], dark, 0.65)
        line(bus_left, cat["bus_y"], bus_right, cat["bus_y"], dark, 0.65)
        for branch in cat["branches"]:
            line(branch["trunk_x"], cat["bus_y"], branch["trunk_x"], branch["last_child_y"], dark, 0.55)
            for node_id in branch["node_ids"]:
                child = node_by_id[node_id]
                fy = child["y"] + child["h"] / 2
                line(branch["trunk_x"], fy, child["x"], fy, dark, 0.45)
    for cat in layout["category_nodes"]:
        rect(cat["x"], cat["y"], cat["w"], cat["h"], cat["colors"]["dark"], radius=3.2)
        lines = cat["cn_lines"] if locale == "cn" else cat["en_lines"]
        text_lines(cat["x"] + cat["w"] / 2, cat["y"] + cat["h"] / 2, lines, category_size, "#FFFFFF", bold, line_height=sizes["category_line"])
    for node in layout["factor_nodes"]:
        rect(node["x"], node["y"], node["w"], node["h"], node["colors"]["light"], node["colors"]["dark"], 2.2, 0.45)
        lines = node["cn_lines"] if locale == "cn" else node["en_lines"]
        text_lines(node["name_x"], node["y"] + node["h"] / 2, lines, factor_size, "#1F2933", regular, "left", sizes["factor_line"])
        text_lines(node["frequency_x"], node["y"] + node["h"] / 2, [f'n={node["frequency"]}'], sizes["frequency"], node["colors"]["dark"], bold, "right")
    note_lines = layout["cn_note_lines"] if locale == "cn" else layout["en_note_lines"]
    text_lines(width / 2, layout["note_y"], note_lines, sizes["note"], "#4B5563", regular, line_height=6.4)
    c.showPage()
    c.save()


def render_png_compact(data: dict[str, Any], layout: dict[str, Any], locale: str, path: Path) -> None:
    """Rasterize the point geometry natively at 600 dpi (not metadata-only)."""
    scale = PNG_DPI / 72.0
    width_px = round(layout["width"] * scale)
    height_px = round(layout["height"] * scale)
    image = Image.new("RGB", (width_px, height_px), "#FFFFFF")
    draw = ImageDraw.Draw(image)
    sizes = layout["sizes"]
    regular_path = pick_font(CN_REGULAR_CANDIDATES if locale == "cn" else EN_REGULAR_CANDIDATES)
    bold_path = pick_font(CN_BOLD_CANDIDATES if locale == "cn" else EN_BOLD_CANDIDATES)

    def s(value: float) -> int:
        return round(value * scale)

    def font(size_pt: float, bold: bool = False) -> ImageFont.FreeTypeFont:
        return ImageFont.truetype(str(bold_path if bold else regular_path), max(1, s(size_pt)))

    def rounded_rect(x: float, y: float, w: float, h: float, fill: str, outline: str | None = None, radius: float = 2.2, line_width: float = 0.45) -> None:
        draw.rounded_rectangle((s(x), s(y), s(x + w), s(y + h)), radius=max(1, s(radius)), fill=fill, outline=outline, width=max(1, s(line_width)) if outline else 1)

    def line(x1: float, y1: float, x2: float, y2: float, fill: str, width_pt: float) -> None:
        draw.line((s(x1), s(y1), s(x2), s(y2)), fill=fill, width=max(1, s(width_pt)))

    def text_lines(x: float, y: float, lines: list[str], size_pt: float, fill: str, bold: bool = False, align: str = "center", line_height: float | None = None) -> None:
        fnt = font(size_pt, bold)
        lh_pt = line_height or size_pt * 1.2
        start_y = y - (len(lines) - 1) * lh_pt / 2
        for i, value in enumerate(lines):
            box = draw.textbbox((0, 0), value, font=fnt)
            tw = box[2] - box[0]
            th = box[3] - box[1]
            if align == "left":
                tx = s(x)
            elif align == "right":
                tx = s(x) - tw
            else:
                tx = s(x) - tw / 2
            center_y = s(start_y + i * lh_pt)
            draw.text((tx, center_y - th / 2 - box[1]), value, font=fnt, fill=fill)

    title_lines = layout["cn_title_lines"] if locale == "cn" else layout["en_title_lines"]
    subtitle_lines = layout["cn_subtitle_lines"] if locale == "cn" else layout["en_subtitle_lines"]
    title_size = sizes["title_cn"] if locale == "cn" else sizes["title_en"]
    category_size = sizes["category_cn"] if locale == "cn" else sizes["category_en"]
    factor_size = sizes["factor_cn"] if locale == "cn" else sizes["factor_en"]
    text_lines(layout["width"] / 2, layout["title_y"], title_lines, title_size, "#1F2933", True, line_height=11.6)
    text_lines(layout["width"] / 2, layout["subtitle_y"], subtitle_lines, sizes["subtitle"], "#4B5563", line_height=7.0)
    rounded_rect(layout["width"] / 2 - 92, layout["legend_y"] - 3.0, 9, 6, "#24557A", radius=1.2)
    text_lines(layout["width"] / 2 - 78, layout["legend_y"], ["一级分类" if locale == "cn" else "Primary category"], sizes["legend"], "#374151", align="left")
    rounded_rect(layout["width"] / 2 + 16, layout["legend_y"] - 3.0, 9, 6, "#DCEAF3", "#24557A", 1.2)
    text_lines(layout["width"] / 2 + 30, layout["legend_y"], ["标准因素" if locale == "cn" else "Standard factor"], sizes["legend"], "#374151", align="left")

    node_by_id = {node["node_id"]: node for node in layout["factor_nodes"]}
    for cat in layout["category_nodes"]:
        if not cat["branches"]:
            continue
        dark = cat["colors"]["dark"]
        trunk_xs = [branch["trunk_x"] for branch in cat["branches"]]
        bus_left = min(trunk_xs + [cat["header_center_x"]])
        bus_right = max(trunk_xs + [cat["header_center_x"]])
        line(cat["header_center_x"], cat["header_bottom_y"], cat["header_center_x"], cat["bus_y"], dark, 0.65)
        line(bus_left, cat["bus_y"], bus_right, cat["bus_y"], dark, 0.65)
        for branch in cat["branches"]:
            line(branch["trunk_x"], cat["bus_y"], branch["trunk_x"], branch["last_child_y"], dark, 0.55)
            for node_id in branch["node_ids"]:
                child = node_by_id[node_id]
                fy = child["y"] + child["h"] / 2
                line(branch["trunk_x"], fy, child["x"], fy, dark, 0.45)
    for cat in layout["category_nodes"]:
        rounded_rect(cat["x"], cat["y"], cat["w"], cat["h"], cat["colors"]["dark"], radius=3.2)
        lines = cat["cn_lines"] if locale == "cn" else cat["en_lines"]
        text_lines(cat["x"] + cat["w"] / 2, cat["y"] + cat["h"] / 2, lines, category_size, "#FFFFFF", True, line_height=sizes["category_line"])
    for node in layout["factor_nodes"]:
        rounded_rect(node["x"], node["y"], node["w"], node["h"], node["colors"]["light"], node["colors"]["dark"], 2.2, 0.45)
        lines = node["cn_lines"] if locale == "cn" else node["en_lines"]
        text_lines(node["name_x"], node["y"] + node["h"] / 2, lines, factor_size, "#1F2933", align="left", line_height=sizes["factor_line"])
        text_lines(node["frequency_x"], node["y"] + node["h"] / 2, [f'n={node["frequency"]}'], sizes["frequency"], node["colors"]["dark"], True, "right")
    note_lines = layout["cn_note_lines"] if locale == "cn" else layout["en_note_lines"]
    text_lines(layout["width"] / 2, layout["note_y"], note_lines, sizes["note"], "#4B5563", line_height=6.4)
    image.save(path, format="PNG", dpi=(PNG_DPI, PNG_DPI), optimize=True)


def create_preview_crops(path: Path, prefix: str) -> None:
    with Image.open(path) as image:
        crop_h = min(1800, image.height)
        starts = [0, max(0, (image.height - crop_h) // 2), max(0, image.height - crop_h)]
        for label, start in zip(["top", "middle", "bottom"], starts):
            crop = image.crop((0, start, image.width, min(image.height, start + crop_h)))
            target_w = 1400
            if crop.width > target_w:
                target_h = max(1, round(crop.height * target_w / crop.width))
                crop = crop.resize((target_w, target_h), Image.Resampling.LANCZOS)
            crop.save(WORK / f"preview_{prefix}_{label}.png", optimize=True)


def prepare() -> dict[str, Any]:
    OUT.mkdir(parents=True, exist_ok=True)
    WORK.mkdir(parents=True, exist_ok=True)
    data = load_and_analyze()
    layout = build_layout(data)
    data["tree_layout"] = {
        "design": layout["design"],
        "dpi": layout["dpi"],
        "width": layout["width"],
        "height": layout["height"],
        "physical_width_in": layout["physical_width_in"],
        "physical_height_in": layout["physical_height_in"],
        "panel_rows": layout["panel_rows"],
        "category_nodes": len(layout["category_nodes"]),
        "factor_nodes": len(layout["factor_nodes"]),
    }
    render_svg_compact(data, layout, "cn", CN_SVG)
    render_pdf_compact(data, layout, "cn", CN_PDF)
    render_png_compact(data, layout, "cn", CN_PNG)
    render_svg_compact(data, layout, "en", EN_SVG)
    render_pdf_compact(data, layout, "en", EN_PDF)
    render_png_compact(data, layout, "en", EN_PNG)
    create_preview_crops(CN_PNG, "cn")
    create_preview_crops(EN_PNG, "en")
    DATA_JSON.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"status": "prepared", "metrics": data["metrics"], "layout": data["tree_layout"]}, ensure_ascii=False, indent=2))
    return data


def svg_node_signature(path: Path) -> dict[str, Any]:
    root = ET.parse(path).getroot()
    ns = {"svg": "http://www.w3.org/2000/svg"}
    categories = []
    factors = []
    for g in root.findall(".//svg:g", ns):
        cls = g.attrib.get("class")
        if cls == "category-node":
            categories.append({k: g.attrib.get(k) for k in ["data-node-id", "data-x", "data-y", "data-w", "data-h"]})
        elif cls == "factor-node":
            factors.append({k: g.attrib.get(k) for k in ["data-node-id", "data-frequency", "data-x", "data-y", "data-w", "data-h"]})
    return {"categories": categories, "factors": factors, "width": root.attrib.get("width"), "height": root.attrib.get("height")}


def validate() -> dict[str, Any]:
    if not DATA_JSON.exists():
        raise FileNotFoundError(f"统计底层数据不存在：{DATA_JSON}")
    data = json.loads(DATA_JSON.read_text(encoding="utf-8"))
    checks: list[dict[str, Any]] = []

    def check(name: str, condition: bool, detail: Any = "") -> None:
        checks.append({"检查项": name, "状态": "通过" if condition else "失败", "详情": detail})
        if not condition:
            raise AssertionError(f"{name}：{detail}")

    metrics = data["metrics"]
    check("全部一级分类总频次之和等于有效记录数", sum(x["total_frequency"] for x in data["category_summary"]) == metrics["valid_records"], metrics["valid_records"])
    check("图示因素频次与singleton数量之和等于有效记录数", metrics["shown_frequency"] + metrics["singleton_factor_count"] == metrics["valid_records"], f"{metrics['shown_frequency']}+{metrics['singleton_factor_count']}={metrics['valid_records']}")
    check("主表与附表完整分割全部标准因素", metrics["ge2_factor_count"] + metrics["singleton_factor_count"] == metrics["standard_factor_count"], metrics["standard_factor_count"])
    main_names = {(r["category_cn"], r["factor_cn"]) for r in data["main_rows"]}
    singleton_names = {(r["category_cn"], r["factor_cn"]) for r in data["singleton_rows"]}
    all_names = {(r["category_cn"], r["factor_cn"]) for r in data["factor_summary"]}
    check("任何标准因素不同时出现在主表和附表", not (main_names & singleton_names), sorted(main_names & singleton_names))
    check("任何标准因素未遗漏", main_names | singleton_names == all_names, len(main_names | singleton_names))
    check("主表仅包含频次≥2因素", all(r["frequency"] >= 2 for r in data["main_rows"]), "")
    check("附表仅包含频次=1因素", all(r["frequency"] == 1 for r in data["singleton_rows"]), "")

    expected_category_order = sorted(data["category_summary"], key=lambda r: (-r["total_frequency"], r["category_en"].casefold()))
    check("一级分类排序符合规则", expected_category_order == data["category_summary"], [r["category_cn"] for r in data["category_summary"]])
    for category in data["category_summary"]:
        rows = [r for r in data["main_rows"] if r["category_cn"] == category["category_cn"]]
        expected = sorted(rows, key=lambda r: (-r["frequency"], r["factor_en"].casefold(), r["factor_cn"]))
        check(f"类内排序符合规则：{category['category_cn']}", rows == expected, len(rows))

    cn_sig = svg_node_signature(CN_SVG)
    en_sig = svg_node_signature(EN_SVG)
    check("中英文SVG节点数量相同", len(cn_sig["categories"]) == len(en_sig["categories"]) and len(cn_sig["factors"]) == len(en_sig["factors"]), f"{len(cn_sig['categories'])}+{len(cn_sig['factors'])}")
    check("中英文SVG节点位置与尺寸完全一致", cn_sig == en_sig, "")
    check("树状图节点总数正确", len(cn_sig["categories"]) + len(cn_sig["factors"]) == metrics["tree_node_count"], metrics["tree_node_count"])
    check("树状图不存在频次1因素", all(int(x["data-frequency"] or 0) >= 2 for x in cn_sig["factors"]), "")

    rebuilt_layout = build_layout(data)
    all_factor_ids = {node["node_id"] for node in rebuilt_layout["factor_nodes"]}
    branch_factor_ids = {
        node_id
        for category in rebuilt_layout["category_nodes"]
        for branch in category["branches"]
        for node_id in branch["node_ids"]
    }
    check(f"{metrics['ge2_factor_count']}个因素节点均且仅均连接到一个一级分类", branch_factor_ids == all_factor_ids and sum(len(branch["node_ids"]) for category in rebuilt_layout["category_nodes"] for branch in category["branches"]) == len(all_factor_ids), len(branch_factor_ids))

    def rects_overlap(a: dict[str, Any], b: dict[str, Any], tolerance: float = 0.05) -> bool:
        return not (
            a["x"] + a["w"] <= b["x"] + tolerance
            or b["x"] + b["w"] <= a["x"] + tolerance
            or a["y"] + a["h"] <= b["y"] + tolerance
            or b["y"] + b["h"] <= a["y"] + tolerance
        )

    factor_rects = [{"x": n["x"], "y": n["y"], "w": n["w"], "h": n["h"], "id": n["node_id"]} for n in rebuilt_layout["factor_nodes"]]
    overlaps = [
        (a["id"], b["id"])
        for idx, a in enumerate(factor_rects)
        for b in factor_rects[idx + 1 :]
        if rects_overlap(a, b)
    ]
    check("紧凑排版中的因素节点不存在重叠", not overlaps, overlaps[:10])
    check(
        "所有节点均位于画布边界内",
        all(n["x"] >= 0 and n["y"] >= 0 and n["x"] + n["w"] <= rebuilt_layout["width"] and n["y"] + n["h"] <= rebuilt_layout["height"] for n in rebuilt_layout["factor_nodes"] + rebuilt_layout["category_nodes"]),
        "",
    )

    for cn_path, en_path, label in [(CN_PNG, EN_PNG, "PNG"), (CN_PDF, EN_PDF, "PDF")]:
        check(f"中英文{label}文件存在且非0字节", cn_path.exists() and en_path.exists() and cn_path.stat().st_size > 0 and en_path.stat().st_size > 0, "")
    with Image.open(CN_PNG) as cn_img, Image.open(EN_PNG) as en_img:
        check("中英文PNG尺寸完全一致", cn_img.size == en_img.size, cn_img.size)
        cn_dpi = cn_img.info.get("dpi", (0, 0))
        en_dpi = en_img.info.get("dpi", (0, 0))
        check("中文PNG分辨率达到600 dpi", min(cn_dpi) >= 599, cn_dpi)
        check("英文PNG分辨率达到600 dpi", min(en_dpi) >= 599, en_dpi)
        expected_size = (
            round(data["tree_layout"]["width"] * PNG_DPI / 72.0),
            round(data["tree_layout"]["height"] * PNG_DPI / 72.0),
        )
        check("PNG像素尺寸与600 dpi原生重绘尺寸一致", cn_img.size == en_img.size == expected_size, expected_size)
        check("中文PNG不是空白图", ImageStat.Stat(cn_img.convert("L")).extrema[0][0] < 250, ImageStat.Stat(cn_img.convert("L")).extrema[0])
        check("英文PNG不是空白图", ImageStat.Stat(en_img.convert("L")).extrema[0][0] < 250, ImageStat.Stat(en_img.convert("L")).extrema[0])
    cn_reader, en_reader = PdfReader(str(CN_PDF)), PdfReader(str(EN_PDF))
    check("中英文PDF均可打开且各含1页", len(cn_reader.pages) == 1 and len(en_reader.pages) == 1, f"{len(cn_reader.pages)}/{len(en_reader.pages)}")
    cn_box, en_box = cn_reader.pages[0].mediabox, en_reader.pages[0].mediabox
    check("中英文PDF页面尺寸完全一致", (float(cn_box.width), float(cn_box.height)) == (float(en_box.width), float(en_box.height)), f"{float(cn_box.width)}x{float(cn_box.height)}")
    check(
        "PDF页面尺寸与论文紧凑版画布一致",
        abs(float(cn_box.width) - data["tree_layout"]["width"]) < 0.1 and abs(float(cn_box.height) - data["tree_layout"]["height"]) < 0.1,
        f"{float(cn_box.width)}x{float(cn_box.height)} pt",
    )
    check(
        "论文紧凑版宽度为7英寸且高宽比不超过1.5",
        abs(data["tree_layout"]["physical_width_in"] - 7.0) < 0.01 and data["tree_layout"]["height"] / data["tree_layout"]["width"] <= 1.5,
        f"{data['tree_layout']['physical_width_in']}x{data['tree_layout']['physical_height_in']} in",
    )

    from openpyxl import load_workbook

    expected_sheets = {
        CN_XLSX: ["树状图", "频次≥2主表", "频次=1附表", "统计说明", "数据一致性问题"],
        EN_XLSX: ["Figure", "Factors_n_GE_2", "Singleton_Appendix", "Statistical_Notes", "Data_Consistency_Issues"],
        BASE_XLSX: ["一级分类汇总", "标准因素频次", "频次1展示分组", "数据问题"],
        ISSUE_XLSX: ["异常汇总", "异常明细", "完全重复数据行"],
    }
    for path, sheets in expected_sheets.items():
        check(f"正式工作簿存在且非0字节：{path.name}", path.exists() and path.stat().st_size > 0, path.stat().st_size if path.exists() else 0)
        wb = load_workbook(path, read_only=False, data_only=False)
        check(f"工作表结构正确：{path.name}", wb.sheetnames == sheets, wb.sheetnames)
        if path in (CN_XLSX, EN_XLSX):
            embedded_images = wb[wb.sheetnames[0]]._images
            check(f"树状图PNG已插入首张工作表：{path.name}", len(embedded_images) == 1, len(embedded_images))
            embedded_hash = hashlib.sha256(embedded_images[0]._data()).hexdigest().upper()
            expected_png = CN_PNG if path == CN_XLSX else EN_PNG
            check(f"工作簿嵌入图已同步为600 dpi紧凑版：{path.name}", embedded_hash == sha256(expected_png), embedded_hash)
        wb.close()

    cn_wb = load_workbook(CN_XLSX, read_only=True, data_only=False)
    en_wb = load_workbook(EN_XLSX, read_only=True, data_only=False)
    cn_main_count = sum(1 for row in cn_wb["频次≥2主表"].iter_rows(min_row=4, values_only=True) if row[5] in {r["factor_cn"] for r in data["main_rows"]})
    en_main_count = sum(1 for row in en_wb["Factors_n_GE_2"].iter_rows(min_row=4, values_only=True) if row[5] in {r["factor_en"] for r in data["main_rows"]})
    cn_single_count = sum(1 for row in cn_wb["频次=1附表"].iter_rows(min_row=4, values_only=True) if row[2] in {r["factor_cn"] for r in data["singleton_rows"]})
    en_single_count = sum(1 for row in en_wb["Singleton_Appendix"].iter_rows(min_row=4, values_only=True) if row[2] in {r["factor_en"] for r in data["singleton_rows"]})
    cn_wb.close(); en_wb.close()
    check(f"中英文主表均完整覆盖{metrics['ge2_factor_count']}个图示因素", cn_main_count == en_main_count == metrics["ge2_factor_count"], f"{cn_main_count}/{en_main_count}")
    check(f"中英文附表均完整覆盖{metrics['singleton_factor_count']}个singleton因素", cn_single_count == en_single_count == metrics["singleton_factor_count"], f"{cn_single_count}/{en_single_count}")

    source_after = file_state(INPUT)
    check("正式输入文件SHA-256未变化", source_after["sha256"] == data["source"]["sha256"], source_after["sha256"])
    check("正式输入文件大小与修改时间未变化", source_after["size"] == data["source"]["size"] and source_after["mtime_ns"] == data["source"]["mtime_ns"], source_after)
    formal_outputs = [CN_SVG, CN_PDF, CN_PNG, EN_SVG, EN_PDF, EN_PNG, CN_XLSX, EN_XLSX, BASE_XLSX, LOG_PATH, ISSUE_XLSX]
    check("所有正式输出均位于visualization目录", all(p.parent.resolve() == OUT.resolve() for p in formal_outputs), "")
    unexpected_artifacts = sorted(
        p.name for p in OUT.iterdir()
        if p.name.startswith("~") or p.name.endswith(".inspect.ndjson")
    )
    check("正式输出目录不存在临时文件或artifact-tool检查副产物", not unexpected_artifacts, unexpected_artifacts)

    if ARTIFACT_VISUAL_JSON.exists():
        artifact_visual = json.loads(ARTIFACT_VISUAL_JSON.read_text(encoding="utf-8"))
        check("四个正式工作簿均通过artifact-tool导入", artifact_visual.get("summary", {}).get("import_failures") == 0, artifact_visual.get("summary"))
        check("全部正式工作表均通过artifact-tool渲染", artifact_visual.get("summary", {}).get("render_failures") == 0 and artifact_visual.get("summary", {}).get("rendered_sheet_count") == 17, artifact_visual.get("summary"))
        check("全部正式工作簿公式错误扫描通过", artifact_visual.get("summary", {}).get("formula_error_workbooks") == 0, artifact_visual.get("summary"))
    else:
        raise FileNotFoundError(f"缺少artifact-tool视觉验证结果：{ARTIFACT_VISUAL_JSON}")

    result = {
        "status": "passed",
        "validated_at": now_iso(),
        "source_before": data["source"],
        "source_after": source_after,
        "metrics": metrics,
        "category_summary": data["category_summary"],
        "checks": checks,
        "outputs": {p.name: {"path": str(p), "size": p.stat().st_size, "sha256": sha256(p)} for p in [CN_SVG, CN_PDF, CN_PNG, EN_SVG, EN_PDF, EN_PNG, CN_XLSX, EN_XLSX, BASE_XLSX, ISSUE_XLSX]},
    }
    VALIDATION_JSON.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    write_log(data, result)
    print(json.dumps({"status": "passed", "checks": len(checks), "metrics": metrics, "report": str(VALIDATION_JSON)}, ensure_ascii=False, indent=2))
    return result


def write_log(data: dict[str, Any], validation: dict[str, Any]) -> None:
    metrics = data["metrics"]
    lines = [
        "# 标准因素两级树状图与双语统计表——可视化运行日志",
        "",
        f"- 运行完成时间：{validation['validated_at']}",
        f"- 输入文件：`{INPUT}`",
        "- 输入工作表：`因素提取`",
        f"- 输入文件SHA-256：`{data['source']['sha256']}`",
        f"- 输入文件大小：{data['source']['size']} 字节",
        f"- 实际使用列：{'、'.join(data['columns'])}",
        f"- 有效记录数：{metrics['valid_records']}",
        f"- 一级分类数：{metrics['category_count']}",
        f"- 结果表实际标准因素数：{metrics['standard_factor_count']}",
        f"- 频次≥2标准因素数：{metrics['ge2_factor_count']}",
        f"- 频次=1标准因素数：{metrics['singleton_factor_count']}",
        f"- 图中节点总数：{metrics['tree_node_count']}",
        f"- 图形排版：论文正文紧凑多面板树图（{data['tree_layout']['physical_width_in']} × {data['tree_layout']['physical_height_in']} 英寸）",
        f"- PNG原生输出：{data['tree_layout']['dpi']} dpi，{round(data['tree_layout']['width'] * data['tree_layout']['dpi'] / 72)} × {round(data['tree_layout']['height'] * data['tree_layout']['dpi'] / 72)} 像素",
        "",
        "## 一级分类统计",
        "",
        "| 排序 | 一级分类 | 总频次 | 图示频次 | 图示因素数 | 频次1因素数 |",
        "|---:|---|---:|---:|---:|---:|",
    ]
    for item in data["category_summary"]:
        lines.append(f"| {item['order']} | {item['category_cn']} | {item['total_frequency']} | {item['shown_frequency']} | {item['ge2_factor_count']} | {item['singleton_factor_count']} |")
    lines.extend(
        [
            "",
            "## 数据一致性检查",
            "",
            f"- 本次共记录 {len(data['issues'])} 项提示或异常；未发现会阻断本次可视化重建的跨分类或中英文一对多冲突。",
        ]
    )
    for issue in data["issues"]:
        lines.append(
            f"- **{issue['问题类型']}**（{issue['严重程度']}，涉及：{issue['涉及对象']}；行号：{issue['涉及行号']}）："
            f"{issue['问题详情']} 处理：{issue['处理方式']}"
        )
    lines.extend(
        [
            "- 详细问题已写入中英文统计表的“数据一致性问题”工作表及 `数据一致性异常.xlsx`。",
            "",
            "## 正式文件",
            "",
        ]
    )
    for name, item in validation["outputs"].items():
        lines.append(f"- `{item['path']}`；{item['size']} 字节；SHA-256 `{item['sha256']}`")
    lines.extend(
        [
            "",
            "## 软件与实现",
            "",
            "- 数据统计：pandas（逐行频次口径）",
            "- SVG：脚本生成的可编辑紧凑多面板节点—连线矢量图",
            "- PDF：ReportLab按7英寸论文正文宽度矢量绘制并嵌入字体",
            "- PNG：Pillow按同一几何原生重绘，600 dpi（非仅修改元数据）",
            "- Excel：@oai/artifact-tool；已同步插入对应语言600 dpi紧凑版PNG；公式、格式、筛选与冻结窗格已验证",
            "",
            "## 最终验证",
            "",
        ]
    )
    for item in validation["checks"]:
        lines.append(f"- [{'x' if item['状态'] == '通过' else ' '}] {item['检查项']} — {item['状态']}" + (f"（{item['详情']}）" if item["详情"] not in ("", None, []) else ""))
    LOG_PATH.write_text("\n".join(lines) + "\n", encoding="utf-8")


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("mode", choices=["prepare", "validate", "all"], nargs="?", default="all")
    args = parser.parse_args()
    if args.mode in ("prepare", "all"):
        prepare()
    if args.mode == "all":
        node = Path(r"C:\Users\LENOVO\.cache\codex-runtimes\codex-primary-runtime\dependencies\node\bin\node.exe")
        builder = WORK / "build_visualization_workbooks.mjs"
        subprocess.run([str(node), str(builder)], check=True, cwd=str(WORK))
        verifier = WORK / "verify_visualization_workbooks.mjs"
        subprocess.run([str(node), str(verifier)], check=True, cwd=str(WORK))
    if args.mode in ("validate", "all"):
        validate()
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        raise
